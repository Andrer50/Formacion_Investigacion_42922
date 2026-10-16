import sys
import io
from docx import Document
import pandas as pd

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

docx_path = 'project/shared/Articulo_RSL_Introduccion_y_Metodologia_IEEE.docx'
doc = Document(docx_path)

print(f"=== AUDITORÍA FINAL DE: {docx_path} ===")
print(f"Total párrafos: {len(doc.paragraphs)}")
print(f"Total tablas: {len(doc.tables)}")

text_dump = " ".join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        text_dump += " " + " ".join([c.text for c in row.cells])

# Check items:
checks = {
    "Early Access Note": "Early Access" in text_dump,
    "Iteración 1 vs Iteración 2": "Iteración 1" in text_dump and "Iteración 2" in text_dump,
    "Consenso y Asesor (Jimmy Sánchez)": "Jimmy Sánchez Portugal" in text_dump and "consenso razonado" in text_dump,
    "Kappa 0.88 y 0.92": "0.88" in text_dump and "0.92" in text_dump,
    "Inviabilidad de Metaanálisis": "metaanálisis" in text_dump and "heterogeneidad métrica" in text_dump,
    "Estudios Nucleares (4)": "[S001]" in text_dump and "[S004]" in text_dump,
    "Sin 'vol. 202x' en referencias": "vol. 202" not in text_dump,
    "Corpus primario [133] y Metodológicas [140]": "[133]" in text_dump and "[140]" in text_dump,
    "Cero 'arxiv' en referencias": "arxiv" not in text_dump.lower(),
    "Nguyen M.-T. (2023, ESWA Q1) como estudio primario [67]": "Gain more with less" in text_dump,
    "TMLR link oficial (sin hashes)": "https://jmlr.org/tmlr/papers/" in text_dump,
    "Métodos formales citados [134]-[140] en Sección II.L": "[134]" in text_dump and "[137]" in text_dump and "[140]" in text_dump
}

for k, v in checks.items():
    print(f"[{'PASS' if v else 'FAIL'}] {k}")

# Check Table VIII (AMSTAR 2)
t_am = doc.tables[7] # Table VIII
print("\n--- Verificación Tabla VIII (Primeras 3 filas y última fila) ---")
for r_idx in [1, 2, 3, len(t_am.rows)-1]:
    row = t_am.rows[r_idx]
    print(f"Fila {r_idx}: Autor='{row.cells[0].text[:35]}' | Base='{row.cells[1].text}' | Cód='{row.cells[2].text}' | Causal='{row.cells[3].text[:50]}...'")

dots_count = sum(1 for row in t_am.rows if "..." in row.cells[3].text or "..." in row.cells[0].text)
print(f"Filas con '...' en Tabla VIII: {dots_count} (Esperado: 0)")

# Check Table IX (Nucleares)
t_nuc = doc.tables[8]
print("\n--- Verificación Tabla IX (Estudios Nucleares y Columna Ref.) ---")
for r in t_nuc.rows:
    print(f"ID={r.cells[0].text:6} | Ref={r.cells[1].text:5} | Autores={r.cells[2].text[:25]:25} | DOI={r.cells[4].text}")

# Check Table X (Muestra de Mapeo)
t_map = doc.tables[9] # Table X
print("\n--- Verificación Tabla X (IDs y Refs en Muestra) ---")
for r in t_map.rows[1:]:
    print(f"ID={r.cells[0].text:6} | Ref={r.cells[1].text:5} | Autores={r.cells[2].text[:22]:22} | PIs={r.cells[3].text}{r.cells[4].text}{r.cells[5].text}{r.cells[6].text} | Total={r.cells[7].text} | Estrato={r.cells[8].text}")

